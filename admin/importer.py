"""Bulk creator import from a spreadsheet.

CSV is the native format — `csv` is stdlib, so it works on any server. `.xlsx`
is read only when openpyxl happens to be installed; if it is not, the uploader
is told to save as CSV rather than being left with a silent failure.

Nothing is written until every row has been checked. A file with one bad row
imports nothing and reports the line, rather than leaving the roster half
updated.

PICTURES IN AN .XLSX. A photo is never *in* a cell — a cell holds text. Excel
stores pictures as floating drawings in xl/media/, anchored to a position on
the sheet, and the cell they cover stays empty. They can still be read, and
the anchor's top-left corner names a row, so a picture sitting on a creator's
row is attached to that creator.

Two consequences worth knowing, because both are visible rather than hidden:
a picture dragged so its top-left corner falls into the row below attaches to
whoever is on that row, and reading pictures needs the workbook loaded fully
rather than in read-only streaming mode, which openpyxl does not expose
drawings in at all. CSV cannot carry an image in any form.
"""

import csv
import io

COLUMNS = ["code", "name", "platform", "handle", "followers",
           "city", "tier", "interest", "note", "active", "photo"]

TIERS = {"nano": "Nano", "micro": "Micro", "mid-tier": "Mid-Tier",
         "mid": "Mid-Tier", "macro": "Macro"}
PLATFORMS = {"instagram": "Instagram", "ig": "Instagram",
             "tiktok": "TikTok", "tt": "TikTok"}

TEMPLATE_ROWS = [
    ["", "Noha Magdy", "Instagram", "noha.mgdi", "697000", "Jeddah", "Macro",
     "Skincare", "example row — delete before importing", "yes",
     "insert the picture on this row (.xlsx only)"],
    ["", "Omnya Elmasry", "TikTok", "omnyaelmasry.0", "618900", "", "Macro",
     "", "leave code blank and the portal assigns one", "yes", ""],
]


def template_csv():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    w.writerows(TEMPLATE_ROWS)
    # BOM so Excel opens it as UTF-8 rather than mangling Arabic city names
    return "﻿" + buf.getvalue()


def template_xlsx():
    """The same template as an .xlsx, which is the only format that can carry
    pictures. Raises RuntimeError when openpyxl is missing, so the caller can
    say so instead of serving a broken file."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError(
            "This server cannot write .xlsx (openpyxl is not installed). "
            "Use the CSV template and attach photos in bulk instead.")
    import io as _io
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Creators"
    ws.append(COLUMNS)
    for row in TEMPLATE_ROWS:
        ws.append(row)
    # Wide enough to drop a picture into, and tall enough that one sits on its
    # own row rather than straddling two.
    ws.column_dimensions["K"].width = 26
    for i in (2, 3):
        ws.row_dimensions[i].height = 60
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_images(ws):
    """{sheet row number: image bytes} for the pictures floating on a sheet.

    The row is the one the picture's top-left corner sits in, which is both the
    only thing an anchor reliably gives and what the eye reads as "this row's
    photo". The first picture on a row wins; a second is ignored rather than
    silently overwriting the first.

    `_data()` and `_from` are openpyxl internals — there is no public API for
    reading embedded images — so every step is guarded. A picture this cannot
    read is skipped, never an exception on the whole import.
    """
    out = {}
    for im in getattr(ws, "_images", []) or []:
        frm = getattr(getattr(im, "anchor", None), "_from", None)
        if frm is None:
            continue          # absolute placement: no row to bind it to
        row = frm.row + 1     # openpyxl counts rows from zero, sheets from one
        if row in out:
            continue
        try:
            payload = im._data()
        except Exception:
            continue
        if payload:
            out[row] = payload
    return out


def _rows_from_xlsx(data):
    """Returns ([(sheet row, cells)], {sheet row: image bytes})."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError(
            "This server cannot read .xlsx (openpyxl is not installed). "
            "Open the file in Excel and use File → Save As → CSV.")
    import tempfile
    import os
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    try:
        os.write(fd, data)
        os.close(fd)
        # Not read_only: that mode streams cells and never loads drawings, so
        # every embedded picture would be invisible. Roster sheets are small.
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        table = [(i, ["" if c is None else str(c).strip() for c in row])
                 for i, row in enumerate(ws.iter_rows(values_only=True), start=1)]
        return table, _sheet_images(ws)
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def _rows_from_csv(data):
    """Returns ([(line number, cells)], {}) — a CSV holds text and nothing else,
    so there are never any images."""
    text = data.decode("utf-8-sig", "replace")
    # Sniff the delimiter: exports from Arabic/European locales use ';'
    sample = text[:2048]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    return [(i, [(c or "").strip() for c in row])
            for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim),
                                    start=1)], {}


def parse(data: bytes, filename: str):
    """Return (rows, errors, photos).

    `rows` are dicts ready for db.upsert_creator, minus the code, which the
    caller assigns; each carries `_row`, its real line in the sheet. `photos`
    maps that same row number to image bytes for an .xlsx with pictures on it,
    and is always empty for a CSV.

    Row numbers are the sheet's own, not a position in the filtered list — an
    error naming "row 14" has to point at row 14 of the file the person is
    looking at, and blank lines in the middle used to shift it.
    """
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        table, photos = _rows_from_xlsx(data)
    else:
        table, photos = _rows_from_csv(data)

    table = [(i, r) for i, r in table if any((c or "").strip() for c in r)]
    if not table:
        return [], ["The file is empty."], {}

    header = [h.strip().lower().replace(" ", "_") for h in table[0][1]]
    if "name" not in header:
        return [], ["No 'name' column found. Download the template and use its headings."], {}

    index = {col: header.index(col) for col in COLUMNS if col in header}
    rows, errors = [], []

    for n, raw in table[1:]:
        def cell(col):
            i = index.get(col)
            return (raw[i].strip() if i is not None and i < len(raw) else "")

        person = cell("name")
        if not person:
            continue  # a blank name is an empty row, not an error

        tier_raw = cell("tier").lower()
        tier = TIERS.get(tier_raw)
        if not tier:
            errors.append("Row " + str(n) + " (" + person + "): tier '"
                          + (cell("tier") or "blank") + "' is not Nano, Micro, Mid-Tier or Macro.")
            continue

        plat_raw = cell("platform").lower()
        platform = PLATFORMS.get(plat_raw)
        if not platform:
            errors.append("Row " + str(n) + " (" + person + "): platform '"
                          + (cell("platform") or "blank") + "' is not Instagram or TikTok.")
            continue

        followers_raw = cell("followers").replace(",", "").replace(" ", "")
        followers = None
        if followers_raw:
            if not followers_raw.isdigit():
                errors.append("Row " + str(n) + " (" + person + "): followers '"
                              + cell("followers") + "' is not a number.")
                continue
            followers = int(followers_raw)

        active_raw = cell("active").lower()
        active = 0 if active_raw in ("no", "false", "0", "hidden") else 1

        rows.append({
            "_row": n,
            "code": cell("code").upper() or None,
            "name": person,
            "handle": cell("handle").lstrip("@") or None,
            "platform": platform,
            "followers": followers,
            "city": cell("city") or None,
            "tier": tier,
            "interest": cell("interest") or None,
            "note": cell("note") or None,
            "active": active,
        })
        # The photo column, if someone typed in it, is a label for where to put
        # the picture — not data. The image itself is matched by row.

    seen = {}
    for r in rows:
        if r["code"]:
            if r["code"] in seen:
                errors.append("Code " + r["code"] + " appears twice in the file.")
            seen[r["code"]] = True

    return rows, errors, photos
